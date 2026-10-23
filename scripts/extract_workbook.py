#!/usr/bin/env python3
"""
Nagpur Police Crime Data Workbook Extraction Script
Extracts, validates, and normalizes all 5 sheets from:
Nagpur_City_Police_Crime_Data_With_Zone_Mapping (1).xlsx
Produces a validated JSON export with full provenance.
"""

import os
import sys
import json
import argparse
from datetime import datetime, timezone
import openpyxl

EXPECTED_SHEETS = [
    'Zone-wise Crime',
    'Police Station-wise',
    'Crime Categories',
    'Summary',
    'Source & Notes'
]

# Reviewed category mapping to SafePath classifier labels
CATEGORY_CLASSIFIER_MAPPING = {
    'ASSAULT ON P.S.': 'Assault',
    'ATT.MURDER': 'Assault',
    'CAUSING DEATH BY NEGLIGENCE': 'Other',
    'CHAIN SNATCHING': 'Robbery',
    'CHEATING': 'Other',
    'CR.BR.OF TRUST': 'Other',
    'CRUELTY BY HUSBAND OR HIS RELATIVES': 'Harassment',
    'DACOITY': 'Robbery',
    'EXTORTION': 'Robbery',
    'H.B.Ts.': 'Theft',
    'HURTS': 'Assault',
    'KIDNAPPING': 'Other',
    'M.V.THEFTS': 'Theft',
    'MOLESTATION': 'Harassment',
    'MURDER': 'Assault',
    'OTHER THEFT': 'Theft',
    'OTHERS': 'Other',
    'PRE.OF DACOITY': 'Suspicious Activity',
    'RAPE': 'Assault',
    'ROBBERY': 'Robbery',
    # Aggregate labels remain unmapped to avoid training or classifying aggregates as single incidents
    'Tot. ROBBERY': None,
    'TOTAL RIOTS': None,
    'TOTAL THEFTS': None
}

def extract_workbook(workbook_path, output_path=None):
    if not os.path.exists(workbook_path):
        raise FileNotFoundError(f"Authoritative workbook not found at: {workbook_path}")

    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    actual_sheets = wb.sheetnames

    for sheet in EXPECTED_SHEETS:
        if sheet not in actual_sheets:
            raise ValueError(f"Required sheet '{sheet}' missing from workbook. Present: {actual_sheets}")

    extraction_time = datetime.now(timezone.utc).isoformat()
    warnings = []
    
    # 1. Source & Notes
    sheet_notes = wb['Source & Notes']
    notes_dict = {}
    for r in range(2, sheet_notes.max_row + 1):
        item = sheet_notes.cell(r, 1).value
        details = sheet_notes.cell(r, 2).value
        if item:
            notes_dict[str(item).strip()] = str(details).strip() if details else ""

    provenance = {
        "sourceWorkbook": os.path.basename(workbook_path),
        "sourceUrl": notes_dict.get('Source', 'https://dashboard.nagpurpolice.in/crime-statistics'),
        "extractionMethod": notes_dict.get('Extraction method', 'Parsed embedded JavaScript data arrays from the HTML response supplied by the user.'),
        "zoneMappingSource": notes_dict.get('Zone mapping source', 'Official Nagpur City Police document listing stations under Zones 3, 4 and 5'),
        "mappingLimitation": notes_dict.get('Mapping limitation', 'Zone 3–5 assignments are taken from official document. Unmatched stations marked unverified.'),
        "importantLimitation": notes_dict.get('Important limitation', 'Station records in embedded source do not include zone field.'),
        "extractionTimestamp": extraction_time,
        "notes": notes_dict
    }

    # 2. Zone-wise Crime
    sheet_zone = wb['Zone-wise Crime']
    zone_rows = []
    for r in range(2, sheet_zone.max_row + 1):
        zone_val = sheet_zone.cell(r, 1).value
        year_val = sheet_zone.cell(r, 2).value
        reported_val = sheet_zone.cell(r, 3).value
        rate_val = sheet_zone.cell(r, 4).value

        if zone_val is None:
            continue

        if not isinstance(year_val, (int, float)) or int(year_val) < 2000:
            raise ValueError(f"Invalid year in Zone-wise Crime row {r}: {year_val}")
        if not isinstance(reported_val, (int, float)) or reported_val < 0:
            raise ValueError(f"Invalid reported crimes in Zone-wise Crime row {r}: {reported_val}")
        if not isinstance(rate_val, (int, float)) or rate_val < 0 or rate_val > 100:
            raise ValueError(f"Invalid detection rate in Zone-wise Crime row {r}: {rate_val}")

        zone_rows.append({
            "zone": str(zone_val).strip(),
            "year": int(year_val),
            "reportedCrimes": int(reported_val),
            "detectionRate": float(rate_val)
        })

    if len(zone_rows) != 10:
        warnings.append(f"Expected 10 zone rows (Zones 1-5 x 2023-2024), found {len(zone_rows)}")

    # 3. Police Station-wise
    sheet_station = wb['Police Station-wise']
    station_rows = []
    unverified_count = 0
    verified_count = 0

    for r in range(2, sheet_station.max_row + 1):
        zone_val = sheet_station.cell(r, 1).value
        name_en = sheet_station.cell(r, 2).value
        name_mr = sheet_station.cell(r, 3).value
        year_val = sheet_station.cell(r, 4).value
        reported_val = sheet_station.cell(r, 5).value
        rate_val = sheet_station.cell(r, 6).value
        manpower_val = sheet_station.cell(r, 7).value or 0
        four_wheelers = sheet_station.cell(r, 8).value or 0
        two_wheelers = sheet_station.cell(r, 9).value or 0

        if name_en is None:
            continue

        zone_str = str(zone_val).strip() if zone_val else "Unverified — Zone 1/2 or check current roster"
        is_unverified = "Unverified" in zone_str

        if is_unverified:
            unverified_count += 1
            mapping_status = "unverified"
        else:
            verified_count += 1
            mapping_status = "verified_document"

        if not isinstance(year_val, (int, float)) or int(year_val) not in [2023, 2024]:
            raise ValueError(f"Invalid year in Station-wise row {r}: {year_val}")
        if not isinstance(reported_val, (int, float)) or reported_val < 0:
            raise ValueError(f"Invalid reported crimes in Station-wise row {r}: {reported_val}")
        if not isinstance(rate_val, (int, float)) or rate_val < 0 or rate_val > 100:
            raise ValueError(f"Invalid detection rate in Station-wise row {r}: {rate_val}")

        station_rows.append({
            "stationName": str(name_en).strip(),
            "stationMarathi": str(name_mr).strip() if name_mr else "",
            "sourceLabel": str(name_en).strip(),
            "zone": zone_str,
            "mappingStatus": mapping_status,
            "year": int(year_val),
            "reportedCrimes": int(reported_val),
            "detectionRate": float(rate_val),
            "manpower": int(manpower_val),
            "fourWheelers": int(four_wheelers),
            "twoWheelers": int(two_wheelers)
        })

    if len(station_rows) != 70:
        warnings.append(f"Expected 70 station rows (35 stations x 2023-2024), found {len(station_rows)}")

    # 4. Crime Categories
    sheet_cat = wb['Crime Categories']
    category_rows = []
    unique_cats = set()

    for r in range(2, sheet_cat.max_row + 1):
        cat_en = sheet_cat.cell(r, 1).value
        cat_mr = sheet_cat.cell(r, 2).value
        year_val = sheet_cat.cell(r, 3).value
        reported_val = sheet_cat.cell(r, 4).value
        detected_val = sheet_cat.cell(r, 5).value
        rate_val = sheet_cat.cell(r, 6).value

        if cat_en is None:
            continue

        cat_clean = str(cat_en).strip()
        unique_cats.add(cat_clean)

        if not isinstance(year_val, (int, float)) or int(year_val) not in [2021, 2022, 2023, 2024]:
            raise ValueError(f"Invalid year in Crime Categories row {r}: {year_val}")
        if not isinstance(reported_val, (int, float)) or reported_val < 0:
            raise ValueError(f"Invalid reported count in Crime Categories row {r}: {reported_val}")
        if not isinstance(detected_val, (int, float)) or detected_val < 0:
            raise ValueError(f"Invalid detected count in Crime Categories row {r}: {detected_val}")
        if not isinstance(rate_val, (int, float)) or rate_val < 0 or rate_val > 100:
            raise ValueError(f"Invalid detection rate in Crime Categories row {r}: {rate_val}")

        classifier_cat = CATEGORY_CLASSIFIER_MAPPING.get(cat_clean, None)
        is_aggregate = cat_clean in ['TOTAL THEFTS', 'TOTAL RIOTS', 'Tot. ROBBERY']

        category_rows.append({
            "sourceLabel": cat_clean,
            "categoryMarathi": str(cat_mr).strip() if cat_mr else "",
            "year": int(year_val),
            "reported": int(reported_val),
            "detected": int(detected_val),
            "detectionRate": float(rate_val),
            "classifierMapping": classifier_cat,
            "isAggregate": is_aggregate
        })

    if len(category_rows) != 92:
        warnings.append(f"Expected 92 category rows (23 categories x 4 years), found {len(category_rows)}")

    # 5. Summary
    sheet_sum = wb['Summary']
    summary_rows = []
    for r in range(2, sheet_sum.max_row + 1):
        year_val = sheet_sum.cell(r, 1).value
        if year_val is None:
            continue
        summary_rows.append({
            "year": int(year_val),
            "zoneReportedTotal": int(sheet_sum.cell(r, 2).value or 0),
            "stationReportedTotal": int(sheet_sum.cell(r, 3).value or 0),
            "categoryReportedTotal": int(sheet_sum.cell(r, 4).value or 0),
            "categoryDetectedTotal": int(sheet_sum.cell(r, 5).value or 0),
            "note": str(sheet_sum.cell(r, 6).value or "")
        })

    output_data = {
        "metadata": {
            "sourceWorkbook": os.path.basename(workbook_path),
            "extractionTimestamp": extraction_time,
            "totalZoneRows": len(zone_rows),
            "totalStationRows": len(station_rows),
            "uniqueStations": len(set(s['stationName'] for s in station_rows)),
            "unverifiedStationYearRows": unverified_count,
            "verifiedStationYearRows": verified_count,
            "totalCategoryRows": len(category_rows),
            "uniqueCategories": len(unique_cats),
            "summaryYears": [s['year'] for s in summary_rows],
            "warnings": warnings
        },
        "provenance": provenance,
        "zones": zone_rows,
        "stations": station_rows,
        "categories": category_rows,
        "summary": summary_rows
    }

    if output_path:
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(output_data, f, indent=2, ensure_ascii=False)
        print(f"[SUCCESS] Extracted workbook data written to: {output_path}")

    return output_data

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Extract and validate Nagpur Police Crime Data workbook")
    parser.add_argument('--input', default='Nagpur_City_Police_Crime_Data_With_Zone_Mapping (1).xlsx', help="Path to input workbook")
    parser.add_argument('--output', default='server/src/data/nagpur_crime_data_extracted.json', help="Path to output JSON")
    args = parser.parse_args()

    try:
        data = extract_workbook(args.input, args.output)
        meta = data["metadata"]
        print(f"Extraction Summary:")
        print(f"  - Zone rows: {meta['totalZoneRows']}")
        print(f"  - Station rows: {meta['totalStationRows']} ({meta['uniqueStations']} unique stations)")
        print(f"  - Unverified station rows: {meta['unverifiedStationYearRows']}")
        print(f"  - Category rows: {meta['totalCategoryRows']} ({meta['uniqueCategories']} unique categories)")
        print(f"  - Summary rows: {len(data['summary'])}")
        if meta['warnings']:
            print(f"  - Warnings: {meta['warnings']}")
    except Exception as e:
        print(f"[ERROR] Extraction failed: {e}", file=sys.stderr)
        sys.exit(1)
